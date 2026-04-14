from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "data" / "raw" / "base_dados_desafio_final_beca.xlsx"
OUTPUT_PATH = ROOT / "data" / "mock_clients.json"


def main() -> None:
    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb["base_de_dados"]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rec = {}
        for h, v in zip(headers, row):
            rec[h] = v.isoformat() if hasattr(v, "isoformat") else v
        rows.append(rec)
    OUTPUT_PATH.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Arquivo gerado: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
